"""Offline tests: run each stage on small fixture files. `pytest -q`"""
import json

import pytest
import shapely
from openpyxl import Workbook, load_workbook
from shapely.geometry import LineString, Point, Polygon

import config as C
import s1_source, s2_age_filter, s3_clean, s4_beach, s5_deliver
from common import city_type, clean_city, clean_province, match_key, norm_level, read_rows, write_rows


@pytest.fixture(autouse=True)
def tmp_data(tmp_path, monkeypatch):
    monkeypatch.setattr(C, "DATA", tmp_path / "data")
    monkeypatch.setattr(C, "OUTPUT", tmp_path / "out")
    monkeypatch.setattr(C, "WORKBOOK", tmp_path / "out" / "wb.xlsx")
    C.DATA.mkdir()


def xlsx(name, header, rows):
    wb = Workbook()
    wb.active.append(header)
    for r in rows:
        wb.active.append(r)
    wb.save(C.DATA / name)


def test_normalizers():
    assert [norm_level(x) for x in ["D-III", "D-IV", "S1", "Sp-1", "S2 Terapan", "Profesi"]] == \
        ["D3", "D4", "S1", "SP1", "S2T", "PROFESI"]
    assert clean_city("Kab.  bandung barat") == "Kabupaten Bandung Barat"
    assert clean_city("KOTA ADM. JAKARTA SELATAN") == "Kota Jakarta Selatan"
    assert clean_city("Kabupaten Administrasi Kepulauan Seribu") == "Kabupaten Kepulauan Seribu"
    assert clean_province("Prov. D.I. Yogyakarta") == "DI Yogyakarta"
    assert clean_province("Prov. D.K.I. Jakarta") == "DKI Jakarta"
    assert match_key("Kab. Batubara") == match_key("Batu Bara") and match_key("Pangkajene Kepulauan") == match_key("Pangkajene dan Kepulauan")
    assert city_type("Kab. Bandung") == "kab" and city_type("Kota Bandung") == "kota" and city_type("Bandung") == ""


def test_typical_age():
    t = s2_age_filter.typical_age
    assert t({"S1", "S2"}) == "18-22; 23+ (postgraduate)" and t({"D3"}) == "18-21" and t({"S2", "S3"}) == "23+"


def test_classify():
    c = s2_age_filter.classify
    assert c({"S1", "S2"}) == "keep" and c({"D3"}) == "keep" and c({"D4"}) == "keep"
    assert c({"S2", "S3"}) == "postgrad-only" and c({"D1", "D2"}) == "non-degree" and c(set()) == "no-data"


def write_sources(n):
    xlsx("src_perguruan_tinggi.xlsx", ["npsn", "nama_pt", "nm_bp", "kelompok", "jln", "Kabupaten_pt", "provinsi_pt"],
         [[f"{i:06d}", f"Univ {i % 3}", "Universitas", "Swasta", "-" if i % 2 else "Jl. Raya 1", "Kab. Badung",
           "Prov. Bali"] for i in range(n)])
    levels = {0: "S2", 1: "D1"}
    xlsx("src_program_studi.xlsx", ["npsn", "nm_jenj_didik"], [[f"{i:06d}", levels.get(i, "S1")] for i in range(n)])


def test_s1_gate_fails_on_short_count():
    write_sources(10)
    with pytest.raises(SystemExit, match="Stage 1 gate FAILED"):
        s1_source.main()


def test_end_to_end(monkeypatch):
    n = C.OFFICIAL_TOTAL
    write_sources(n)
    s1_source.main()
    s2_age_filter.main()
    kept = {r["npsn"]: r for r in read_rows("s2_institutions.csv")}
    assert len(kept) == n                                                # AGE_FILTER off: nothing dropped
    assert kept["000000"]["offers_undergrad"] == "no" and kept["000000"]["typical_age"] == "23+"
    assert kept["000001"]["typical_age"] == "18-19" and kept["000002"]["typical_age"] == "18-22"
    s3_clean.main()
    rows = read_rows("s3_institutions.csv")
    assert {r["city"] for r in rows} == {"Kabupaten Badung"} and rows[0]["province"] == "Bali"
    assert all(r["duplicate"] == "yes" for r in rows)                    # names repeat in same city
    assert {r["address"] for r in rows} == {"Jl. Raya 1", C.ADDRESS_MISSING}

    square = Polygon([(0, 0), (1, 0), (1, 1), (0, 1)]).difference(Polygon([(.4, .4), (.6, .4), (.6, .6), (.4, .6)]))
    feats = [
        {"kind": "admin", "id": 1, "tags": {"admin_level": "4", "name": "Bali"}, "geom": Polygon([(-1, -1), (2, -1), (2, 2), (-1, 2)])},
        {"kind": "admin", "id": 77, "tags": {"admin_level": "5", "name": "Kabupaten Badung"}, "geom": square},
        {"kind": "coast", "geom": LineString([(0.1, 0), (0.1, 1)])},
        {"kind": "coast", "geom": LineString([(0.5, 0.45), (0.5, 0.55)])},   # inside the enclave hole
        {"kind": "beach", "osm": ["way", 5], "tags": {"name": "Pantai Kuta", "surface": "sand"}, "geom": Point(0.1005, 0.2)},
        {"kind": "beach", "osm": ["way", 6], "tags": {"name": "Mud Flat", "surface": "mud"}, "geom": Point(0.1005, 0.3)},
        {"kind": "beach", "osm": ["node", 7], "tags": {"name": "Lake Beach"}, "geom": Point(0.9, 0.9)},
        {"kind": "beach", "osm": ["node", 8], "tags": {"name": "Pantai In Hole", "tourism": "attraction"}, "geom": Point(0.5, 0.5)},
    ]
    with open(C.DATA / s4_beach.FEATURES, "w") as f:
        for r in feats:
            f.write(json.dumps({**{k: v for k, v in r.items() if k != "geom"}, "wkb": shapely.to_wkb(r["geom"], hex=True)}) + "\n")
    monkeypatch.setattr(C, "SPOTCHECK_FILE", C.DATA / "spot.csv")
    write_rows("spot.csv", [{"city": "Kabupaten Badung", "province": "Bali", "web_beach": "yes", "evidence": "x"}],
               ["city", "province", "web_beach", "evidence"])
    monkeypatch.setattr(C, "SPOTCHECK_MIN_AGREE", 1)                     # fixture has 1 city, not 20
    s4_beach.main()
    city = read_rows("s4_cities.csv")[0]
    assert city["beach"] == "yes" and city["beach_names"] == "Pantai Kuta" and city["beach_source"].endswith("/way/5")
    spot = read_rows("s4_spotcheck.csv")[0]
    assert spot["agree"] == "yes"

    monkeypatch.setattr(C, "SHORTLIST_NUMERIC_FILE", C.DATA / "num.csv")
    monkeypatch.setattr(C, "SHORTLIST_SAFETY_FILE", C.DATA / "saf.csv")
    write_rows("num.csv", [{"city": "Kabupaten Badung", "province": "Bali", "poverty_pct": "3.5", "poverty_period": "2025",
                            "poverty_src": "https://x/p", "unemployment_pct": "", "unemployment_period": "",
                            "unemployment_src": "UNVERIFIED", "umk_2026_idr": "3500000", "umk_src": "https://x/u",
                            "avg_wage_idr": "not published at regency level", "avg_wage_src": "UNVERIFIED",
                            "notes": ""}],
               ["city", "province", "poverty_pct", "poverty_period", "poverty_src", "unemployment_pct",
                "unemployment_period", "unemployment_src", "umk_2026_idr", "umk_src", "avg_wage_idr", "avg_wage_src",
                "notes"])
    write_rows("saf.csv", [], ["city", "province", "safety", "n_sources", "evidence"])
    monkeypatch.setattr(C, "SHORTLIST_SWIM_FILE", C.DATA / "swim.csv")
    write_rows("swim.csv", [{"city": "Kabupaten Badung", "province": "Bali", "swim_sand_beach": "yes",
                             "beach_name": "Pantai Kuta", "evidence": "sand, swimming - https://x/s"}],
               ["city", "province", "swim_sand_beach", "beach_name", "evidence"])
    s5_deliver.main()
    wb = load_workbook(C.WORKBOOK)
    assert wb.sheetnames == ["Institutions", "Summary", C.SHORTLIST_SHEET]
    short = list(wb[C.SHORTLIST_SHEET].iter_rows(min_row=2, values_only=True))[0]
    assert short[4] == C.NOT_RESEARCHED and short[6] == 3500000 and short[9] == C.UNVERIFIED and short[11] == 3.5
    assert short[8] == C.NOT_AT_REGENCY and short[15] == "Yes" and short[16].startswith("Pantai Kuta: ")
    assert wb["Institutions"].max_row - 1 == len(rows)
    assert list(wb["Summary"].iter_rows(min_row=2, values_only=True)) == [("Kabupaten Badung", "Bali", len(rows), "yes", len(rows) - 2)]


def test_is_beach():
    assert s4_beach.is_beach({"natural": "beach", "name": "Kuta"})
    assert s4_beach.is_beach({"tourism": "attraction", "name": "Pantai Parangtritis"})
    assert not s4_beach.is_beach({"tourism": "hotel", "name": "Pantai Inn"})
    assert not s4_beach.is_beach({"tourism": "attraction", "name": "Danau Toba"})
    assert not s4_beach.is_beach({"natural": "beach"})


def test_s4_gate_fails_on_unmatched_city():
    write_rows("s3_institutions.csv", [{"name": "A", "city": "Kota Nowhere", "province": "Bali"}], ["name", "city", "province"])
    (C.DATA / s4_beach.FEATURES).write_text("")
    with pytest.raises(SystemExit, match="Stage 4 gate FAILED"):
        s4_beach.main()


def test_s5_gate_fails_on_missing_beach():
    row = {"npsn": "1", "name": "A", "type": "Universitas", "ownership": "Private", "address": "x", "city": "Kota Y",
           "province": "Bali", "levels": "S1", "offers_undergrad": "yes", "typical_age": "18-22", "duplicate": "no",
           "city_basis": "official"}
    write_rows("s3_institutions.csv", [row], list(row))
    write_rows("s4_cities.csv", [], ["city", "province", "beach"])
    with pytest.raises(SystemExit, match="Stage 5 gate FAILED"):
        s5_deliver.main()
    assert not C.WORKBOOK.exists()
