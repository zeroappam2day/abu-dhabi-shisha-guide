"""Stage 5: write the 2-sheet workbook and verify it by reading it back."""
from collections import Counter

from openpyxl import Workbook, load_workbook

import config as C
from common import gate, read_rows

INST_COLS = ["NPSN", "Name", "Type", "Ownership", "Address", "City/Regency", "Province", "Levels offered",
             "Offers S1/D3/D4", "Typical student age", "Beach", "Duplicate", "City source"]
SUM_COLS = ["City/Regency", "Province", "Institution count", "Beach", "Offering S1/D3/D4"]


def main():
    rows = read_rows("s3_institutions.csv")
    beach = {(c["city"], c["province"]): c["beach"] for c in read_rows("s4_cities.csv")}
    counts = Counter((r["city"], r["province"]) for r in rows)
    under = Counter((r["city"], r["province"]) for r in rows if r["offers_undergrad"] == "yes")

    wb = Workbook()
    ws = wb.active
    ws.title = "Institutions"
    ws.append(INST_COLS)
    for r in rows:
        ws.append([r["npsn"], r["name"], r["type"], r["ownership"], r["address"], r["city"], r["province"],
                   r["levels"], r["offers_undergrad"], r["typical_age"], beach.get((r["city"], r["province"]), ""),
                   r["duplicate"], r["city_basis"]])
    ws2 = wb.create_sheet("Summary")
    ws2.append(SUM_COLS)
    for (city, prov), n in sorted(counts.items()):
        ws2.append([city, prov, n, beach.get((city, prov), ""), under[(city, prov)]])
    for sheet in (ws, ws2):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    C.OUTPUT.mkdir(exist_ok=True)
    draft = C.WORKBOOK.with_suffix(".unverified.xlsx")  # renamed only after the gate passes
    wb.save(draft)

    back = load_workbook(draft, read_only=True)
    inst = list(back["Institutions"].iter_rows(min_row=2, values_only=True))
    summ = list(back["Summary"].iter_rows(min_row=2, values_only=True))
    empty = [i for i, row in enumerate(inst + summ) if any(v in (None, "") for v in row)]
    gate(5, [
        ("Institutions rows equal Stage 3", len(inst) == len(rows), f"{len(inst)} vs {len(rows)}"),
        ("Summary counts sum to the same total", sum(r[2] for r in summ) == len(rows),
         f"{sum(r[2] for r in summ)} across {len(summ)} cities"),
        ("no empty cells", not empty, f"{len(empty)} rows with empty cells"),
    ])
    back.close()
    draft.replace(C.WORKBOOK)
    print(f"Saved {C.WORKBOOK}")


if __name__ == "__main__":
    main()
