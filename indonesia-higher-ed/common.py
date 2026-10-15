"""Shared helpers: gates, stage I/O, downloads, name normalization."""
import csv
import re
import sys

import config as C


def gate(stage, checks):
    """checks: list of (description, passed, detail). Writes a report; exits 1 on any failure."""
    C.DATA.mkdir(exist_ok=True)
    lines = [f"{'PASS' if ok else 'FAIL'}  {desc}  [{detail}]" for desc, ok, detail in checks]
    report = f"Stage {stage} gate\n" + "\n".join(lines) + "\n"
    (C.DATA / f"gate_s{stage}.txt").write_text(report, encoding="utf-8")
    print(report)
    if not all(ok for _, ok, _ in checks):
        sys.exit(f"Stage {stage} gate FAILED. Stop: report the FAIL lines above and decide before continuing.")


def write_rows(name, rows, fields):
    C.DATA.mkdir(exist_ok=True)
    with open(C.DATA / name, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def read_rows(name):
    path = C.DATA / name
    if not path.exists():
        sys.exit(f"{path} missing: run the previous stage first.")
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def download(url, name):
    """Fetch url into DATA/name once; later runs reuse the file."""
    import requests
    path = C.DATA / name
    if not path.exists():
        C.DATA.mkdir(exist_ok=True)
        with requests.get(url, stream=True, timeout=600) as r:
            r.raise_for_status()
            tmp = path.with_suffix(".part")
            with open(tmp, "wb") as f:
                for chunk in r.iter_content(1 << 20):
                    f.write(chunk)
            tmp.replace(path)
    return path


def read_xlsx(path, cols):
    """Rows of the first sheet as dicts, renamed via cols {ours: theirs}; all values as stripped strings."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    missing = [c for c in cols.values() if c not in header]
    if missing:
        sys.exit(f"{path.name}: columns {missing} not found; header is {header}")
    idx = {ours: header.index(theirs) for ours, theirs in cols.items()}
    return [{k: squash(str(row[i]) if row[i] is not None else "") for k, i in idx.items()} for row in it]


def squash(s):
    return re.sub(r"\s+", " ", (s or "").strip())


def norm_level(raw):
    key = re.sub(r"[^A-Z0-9]", "", (raw or "").upper())
    return C.LEVEL_ALIASES.get(key, key)


def clean_city(raw):
    s = squash(raw)
    for pattern, repl in C.CITY_PREFIXES:
        if re.match(pattern, s, flags=re.I):
            s = re.sub(pattern, repl, s, count=1, flags=re.I)
            break
    head, _, rest = s.partition(" ")
    if head in ("Kota", "Kabupaten"):
        s = f"{head} {rest.title()}"
    else:
        s = s.title()
    return C.CITY_ALIASES.get(s, s)


def clean_province(raw):
    s = squash(raw)
    for pattern, repl in C.PROVINCE_PREFIXES:
        s = re.sub(pattern, repl, s, count=1, flags=re.I)
    return C.PROVINCE_ALIASES.get(s.lower(), s.title())


def city_type(name):
    """'kota', 'kab' or '' (no prefix) for a city name."""
    head = clean_city(name).split(" ", 1)[0].lower()
    return {"kota": "kota", "kabupaten": "kab"}.get(head, "")


def match_key(name):
    """Bare comparison key: prefix, 'dan', spaces and punctuation removed (Kab. Batubara == Batu Bara)."""
    s = re.sub(r"^(kota|kabupaten) ", "", clean_city(name).lower())
    return re.sub(r"[^a-z0-9]", "", re.sub(r"\bdan\b", "", s))
