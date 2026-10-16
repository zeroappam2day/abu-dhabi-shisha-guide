"""Z5: build zone/output/priority_zone.xlsx + .html (original 18 columns unchanged + v2 overlay) and gate."""
import json
import subprocess
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

import zone_config as Z
from zone_common import baseline_ok, gate, read_csv

sys.path.insert(0, str(Z.ROOT))
import build_html  # read-only import: reuse TEMPLATE

V2 = ["Swim (v2)", "Swim evidence (v2)", "Safety (v2)", "Safety evidence (v2)", "Unemployment % (v2)",
      "Poverty % (v2)", "UMK note (v2)", "Changed", "Poverty band (v2)", "Unemployment band (v2)"]


def band_label(lo, hi, name):
    rng = f"<{hi}%" if lo is None else (f">={lo}%" if hi is None else f"{lo}-{hi}%")
    return f"{name} ({rng})"


def band(v, bands):
    if not isinstance(v, (int, float)):
        return Z.BAND_MISSING
    for lo, hi, name in bands:
        if (lo is None or v >= lo) and (hi is None or v < hi):
            return band_label(lo, hi, name)
    raise ValueError(f"{v} outside all bands")

wb = load_workbook(Z.WORKBOOK, read_only=True)
sl = list(wb[Z.SHEET].iter_rows(values_only=True))
cols = list(sl[0])
orig = [r for r in sl[1:] if r[cols.index("Priority zone")] != Z.OUTSIDE]
key = lambda c, p: (c, p)
swim = {key(r["city"], r["province"]): r for r in read_csv(Z.SWIM_V2)}
safe = {key(r["city"], r["province"]): r for r in read_csv(Z.SAFETY_V2)}
num = defaultdict(dict)
for r in read_csv(Z.NUMBERS_V2):
    num[key(r["city"], r["province"])][r["metric"]] = r


def v2row(o):
    d = dict(zip(cols, o))
    k = key(d["City/Regency"], d["Province"])
    s, f, n = swim[k], safe[k], num.get(k, {})
    une, pov = d["Unemployment rate % (all ages)"], d["Poverty rate %"]
    if n.get("unemployment", {}).get("apply") == "yes":
        une = float(n["unemployment"]["value_v2"])
    if n.get("poverty", {}).get("apply") == "yes":
        pov = float(n["poverty"]["value_v2"])
    umk = n.get("umk")
    umk_note = (f"{umk['value_v2']}: {umk['evidence']}" if umk else "")
    changed = [lbl for lbl, new, old in [("Swim", s["swim_v2"], d["Swimmable sand beach"]),
                                         ("Safety", f["safety_v2"], d["Safety (walking around)"]),
                                         ("Unemployment", une, d["Unemployment rate % (all ages)"]),
                                         ("Poverty", pov, d["Poverty rate %"])] if new != old]
    return [s["swim_v2"], s["evidence"], f["safety_v2"], f["evidence"], une, pov, umk_note,
            "; ".join(changed) or "No", band(pov, Z.POVERTY_BANDS), band(une, Z.UNEMPLOYMENT_BANDS)]


rows = [list(o) + v2row(o) for o in orig]
Z.OUTPUT.mkdir(exist_ok=True)
out = Workbook()
ws = out.active
ws.title = "Priority zone"
ws.append(cols + V2)
for r in rows:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True)
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions
out.save(Z.OUT_XLSX)

# HTML: same template and filters; displayed Safety/Swim/Unemployment/Poverty are v2, originals shown in detail.
inst = defaultdict(list)
it = wb["Institutions"].iter_rows(values_only=True)
h = next(it)
ix = {k: h.index(k) for k in h}
for r in it:
    inst[(r[ix["City/Regency"]], r[ix["Province"]])].append(
        [r[ix["Name"]], r[ix["Type"]], r[ix["Ownership"]], r[ix["Levels offered"]], r[ix["Offers S1/D3/D4"]]])
data = []
for r in rows:
    d = dict(zip(cols + V2, r))
    o = {f"orig_{k}": d[k] for k in ["Swimmable sand beach", "Safety (walking around)", "Unemployment rate % (all ages)", "Poverty rate %"]}
    d.update(o)
    d["Swimmable sand beach"], d["Swim beach evidence"] = d["Swim (v2)"], d["Swim evidence (v2)"]
    d["Safety (walking around)"], d["Safety evidence"] = d["Safety (v2)"], d["Safety evidence (v2)"]
    d["Unemployment rate % (all ages)"], d["Poverty rate %"] = d["Unemployment % (v2)"], d["Poverty % (v2)"]
    d["Poverty band"], d["Unemployment band"] = d["Poverty band (v2)"], d["Unemployment band (v2)"]
    d["institutions"] = sorted(inst[(d["City/Regency"], d["Province"])])
    data.append(d)
patches = [
    ("<title>Indonesia Beach Places</title>", "<title>Priority Zone Places</title>"),
    ("<h1>Indonesia beach places for students (17–25)</h1>", "<h1>Priority-zone beach places (Java, Bali, Lampung), v2 evidence</h1>"),
    ('<p class="sub">260 cities/regencies with a real sea beach. Click a column to sort, a row to expand. Safety: 60 largest researched; 200 not researched.',
     '<p class="sub">70 priority-zone places. Safety, swim, unemployment and poverty show re-verified (v2) values; the original values are in each row\'s detail. Click a column to sort, a row to expand.'),
    ('[["260","beach places"]', '[[DATA.length,"places"],[DATA.filter(d=>d.Changed!=="No").length,"changed in v2"]'),
    ('<div class="grid"><div class="box"><h3>Safety evidence</h3>',
     '<div class="grid"><div class="box"><h3>Original values (baseline)</h3><p>Changed: ${esc(d.Changed)}</p><p>Safety: ${esc(d["orig_Safety (walking around)"])}</p><p>Swim: ${esc(d["orig_Swimmable sand beach"])}</p><p>Unemployment %: ${esc(d["orig_Unemployment rate % (all ages)"])}</p><p>Poverty %: ${esc(d["orig_Poverty rate %"])}</p>${d["UMK note (v2)"]?`<p>UMK note (v2): ${linkify(d["UMK note (v2)"])}</p>`:""}</div><div class="box"><h3>Safety evidence (v2)</h3>'),
    ('<details class="ms" id="zone"></details>', '<details class="ms" id="zone"></details><details class="ms" id="pov"></details><details class="ms" id="une"></details>'),
    ("const OFF={prov:new Set(),saf:new Set(),umk:new Set(),swim:new Set(),zone:new Set()}",
     "const OFF={prov:new Set(),saf:new Set(),umk:new Set(),swim:new Set(),zone:new Set(),pov:new Set(),une:new Set()}"),
    ('vals.map(v=>`<label><input type="checkbox" value="${esc(v)}" checked> ${esc(v)}</label>`)',
     'vals.map(v=>`<label><input type="checkbox" value="${esc(v)}" checked> ${esc(v)} <span class="muted cnt" data-v="${esc(v)}"></span></label>`)'),
    ('buildMS("swim",', 'buildMS("pov","Poverty",BANDS.pov.filter(v=>DATA.some(d=>d["Poverty band"]===v)));buildMS("une","Unemployment",BANDS.une.filter(v=>DATA.some(d=>d["Unemployment band"]===v)));\nbuildMS("swim",'),
    ('!OFF.prov.has(d.Province)&&!OFF.saf.has(d["Safety (walking around)"])&&!OFF.umk.has(d["UMK basis"])&&!OFF.swim.has(d["Swimmable sand beach"])&&!OFF.zone.has(d["Priority zone"])&&\n(!q||(d["City/Regency"]+" "+d.Province+" "+d.institutions.map(i=>i[0]).join(" ")).toLowerCase().includes(q)))',
     'pass(d,q))'),
    ("function rows(){", 'const FIELD={prov:"Province",saf:"Safety (walking around)",umk:"UMK basis",swim:"Swimmable sand beach",zone:"Priority zone",pov:"Poverty band",une:"Unemployment band"};\n'
     'function pass(d,q,skip){for(const id in FIELD)if(id!==skip&&OFF[id].has(d[FIELD[id]]))return false;\n'
     'return !q||(d["City/Regency"]+" "+d.Province+" "+d.institutions.map(i=>i[0]).join(" ")).toLowerCase().includes(q)}\n'
     'function counts(){const q=$("q").value.trim().toLowerCase();for(const id in MSVALS){const c={};DATA.forEach(d=>{if(pass(d,q,id))c[d[FIELD[id]]]=(c[d[FIELD[id]]]||0)+1});\n'
     '$(id).querySelectorAll(".cnt").forEach(s=>s.textContent="("+(c[s.dataset.v]||0)+")")}}\n'
     "function rows(){"),
    ("function render(){const r=rows();", "function render(){const r=rows();counts();"),
    ('<td class="num">${fmt(d["Unemployment rate % (all ages)"],"%")}</td>',
     '<td class="num">${fmt(d["Unemployment rate % (all ages)"],"%")}<br><span class="muted bnd">${esc(d["Unemployment band"])}</span></td>'),
    (':fmt(d["Poverty rate %"],"%")}</td>', ':fmt(d["Poverty rate %"],"%")}<br><span class="muted bnd">${esc(d["Poverty band"])}</span></td>'),
    (".muted{color:var(--muted)}", ".muted{color:var(--muted)}.bnd{font-size:11px;white-space:nowrap}"),
]
html = build_html.TEMPLATE
for a, b in patches:
    assert html.count(a) == 1, a[:60]
    html = html.replace(a, b)
bands = {"pov": [band_label(*b) for b in Z.POVERTY_BANDS] + [Z.BAND_MISSING],
         "une": [band_label(*b) for b in Z.UNEMPLOYMENT_BANDS] + [Z.BAND_MISSING]}
html = html.replace("const DATA=__DATA__;", "const BANDS=" + json.dumps(bands) + ";\nconst DATA=__DATA__;")
Z.OUT_HTML.write_text(html.replace("__DATA__", json.dumps(data, ensure_ascii=False).replace("</", "<\\/")), encoding="utf-8")

# Read back and gate
back = list(load_workbook(Z.OUT_XLSX, read_only=True).active.iter_rows(values_only=True))
same_orig = all(tuple(b[:18]) == tuple(o) for b, o in zip(back[1:], orig))
diff = subprocess.run(["git", "diff", "--stat", Z.BASELINE_COMMIT, "--", "indonesia-higher-ed", ":!indonesia-higher-ed/zone"],
                      cwd=Z.ROOT.parent, capture_output=True, text=True).stdout.strip()
html_back = Z.OUT_HTML.read_text(encoding="utf-8")
hdata = json.loads(html_back.split("const DATA=", 1)[1].split(";\nconst COLS", 1)[0].replace("<\\/", "</"))
ok, msg = baseline_ok()
gate("z5", [
    ("xlsx rows", len(back) - 1 == Z.EXPECTED_ROWS, f"{len(back) - 1}"),
    ("xlsx columns", list(back[0]) == cols + V2, f"{len(back[0])} = 18 + {len(V2)}"),
    ("original 18 columns unchanged", same_orig, ""),
    ("no empty v2 cells (except UMK note)", all(all(c not in (None, "") for c in r[18:24] + r[25:]) for r in back[1:]), ""),
    ("html has 70 places with v2 values", len(hdata) == Z.EXPECTED_ROWS and "__DATA__" not in html_back
     and all(d["Safety (walking around)"] == d["Safety (v2)"] and d["Swimmable sand beach"] == d["Swim (v2)"] for d in hdata),
     f"{len(hdata)} places"),
    ("every poverty/unemployment value in exactly one band", all(r[26] and r[27] for r in back[1:]),
     f"pov {dict(Counter(r[26] for r in back[1:]))}; une {dict(Counter(r[27] for r in back[1:]))}"),
    ("baseline intact", ok, msg),
    ("no diff outside zone/ vs e34937f", diff == "", diff[:200] or "empty"),
])
