"""Build output/beach_places.html: a self-contained, sortable, filterable, expandable table of the Beach shortlist."""
import json
from collections import defaultdict

from openpyxl import load_workbook

import config as C

OUT = C.OUTPUT / "beach_places.html"


def main():
    wb = load_workbook(C.WORKBOOK, read_only=True)
    inst = defaultdict(list)
    rows = wb["Institutions"].iter_rows(values_only=True)
    h = next(rows)
    ix = {k: h.index(k) for k in h}
    for r in rows:
        inst[(r[ix["City/Regency"]], r[ix["Province"]])].append(
            [r[ix["Name"]], r[ix["Type"]], r[ix["Ownership"]], r[ix["Levels offered"]], r[ix["Offers S1/D3/D4"]]])
    sl = wb["Beach shortlist"].iter_rows(values_only=True)
    cols = next(sl)
    data = []
    for r in sl:
        d = dict(zip(cols, r))
        d["institutions"] = sorted(inst[(d["City/Regency"], d["Province"])])
        data.append(d)
    payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    OUT.write_text(TEMPLATE.replace("__DATA__", payload), encoding="utf-8")
    print(f"Saved {OUT} ({len(data)} places)")


TEMPLATE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Indonesia Beach Places</title>
<style>
:root{--bg:#f7f7f5;--panel:#fff;--ink:#1d1d1b;--muted:#6b6b66;--line:#e3e2dc;--accent:#1f6f8b;--accent-soft:#e3f0f4;
--low:#b4442c;--low-bg:#fbe9e4;--med:#9a6b00;--med-bg:#fbf1d9;--high:#2e7d4f;--high-bg:#e3f3e9;--none:#77766f;--none-bg:#efeee9}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#141413;--panel:#1d1d1b;--ink:#ecebe6;--muted:#a3a29b;--line:#33322f;
--accent:#6fb6cf;--accent-soft:#1f3640;--low:#ec8a72;--low-bg:#3a201a;--med:#e3b65a;--med-bg:#382d14;--high:#7fcf9e;--high-bg:#183224;--none:#a3a29b;--none-bg:#2a2a27}}
:root[data-theme="dark"]{--bg:#141413;--panel:#1d1d1b;--ink:#ecebe6;--muted:#a3a29b;--line:#33322f;--accent:#6fb6cf;--accent-soft:#1f3640;
--low:#ec8a72;--low-bg:#3a201a;--med:#e3b65a;--med-bg:#382d14;--high:#7fcf9e;--high-bg:#183224;--none:#a3a29b;--none-bg:#2a2a27}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:14px/1.45 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif}
header{padding:20px 16px 8px;max-width:1400px;margin:auto}h1{font-size:20px;margin:0 0 4px}.sub{color:var(--muted);margin:0}
.stats{display:flex;gap:10px;flex-wrap:wrap;margin:14px 0 4px}.stat{background:var(--panel);border:1px solid var(--line);border-radius:8px;padding:8px 12px}
.stat b{display:block;font-size:18px}.stat span{color:var(--muted);font-size:12px}
.controls{position:sticky;top:0;z-index:5;background:var(--bg);padding:10px 16px;max-width:1400px;margin:auto;display:flex;gap:8px;flex-wrap:wrap;align-items:center;border-bottom:1px solid var(--line)}
input,select,button{font:inherit;color:var(--ink);background:var(--panel);border:1px solid var(--line);border-radius:6px;padding:6px 10px}
input[type=search]{flex:1 1 220px;min-width:0}button{cursor:pointer}button:hover{border-color:var(--accent)}
.count{color:var(--muted);margin-left:auto}
main{max-width:1400px;margin:auto;padding:8px 16px 40px;overflow-x:auto}
table{width:100%;border-collapse:collapse;background:var(--panel);border:1px solid var(--line);border-radius:8px;overflow:hidden;min-width:880px}
th{position:sticky;top:0;background:var(--panel);text-align:left;font-weight:600;font-size:12px;color:var(--muted);padding:10px;border-bottom:1px solid var(--line);cursor:pointer;user-select:none;white-space:nowrap}
th:hover{color:var(--ink)}th.sorted{color:var(--accent)}th .arrow{font-size:10px;margin-left:4px}
td{padding:9px 10px;border-bottom:1px solid var(--line);vertical-align:top}td.num{text-align:right;font-variant-numeric:tabular-nums}
tr.row{cursor:pointer}tr.row:hover td{background:var(--accent-soft)}tr.row td:first-child{white-space:nowrap}
.ms{position:relative}.ms summary{list-style:none;cursor:pointer;background:var(--panel);border:1px solid var(--line);border-radius:6px;padding:6px 28px 6px 10px;white-space:nowrap;position:relative}
.ms summary::-webkit-details-marker{display:none}.ms summary::after{content:"▾";position:absolute;right:10px;color:var(--muted)}
.ms.active summary{border-color:var(--accent);color:var(--accent)}
.ms .menu{position:absolute;z-index:20;top:calc(100% + 4px);left:0;min-width:240px;max-height:320px;overflow:auto;background:var(--panel);border:1px solid var(--line);border-radius:8px;padding:6px;box-shadow:0 6px 24px rgba(0,0,0,.15)}
.ms label{display:flex;gap:8px;align-items:center;padding:5px 6px;border-radius:4px;cursor:pointer;white-space:nowrap}.ms label:hover{background:var(--accent-soft)}
.mbtn{display:flex;gap:6px;padding:2px 2px 6px;border-bottom:1px solid var(--line);margin-bottom:4px}.mbtn button{padding:3px 10px;font-size:12px}
.copy{margin-left:6px;padding:2px 4px;border:1px solid transparent;background:none;color:var(--muted);vertical-align:-2px;line-height:0;border-radius:4px}
.copy:hover{color:var(--accent);border-color:var(--line)}.copy.ok{color:var(--high)}
.caret{display:inline-block;width:14px;color:var(--muted);transition:transform .15s}tr.open .caret{transform:rotate(90deg)}
.badge{display:inline-block;padding:2px 8px;border-radius:999px;font-size:12px;white-space:nowrap}
.b-Low{background:var(--low-bg);color:var(--low)}.b-Medium{background:var(--med-bg);color:var(--med)}.b-High{background:var(--high-bg);color:var(--high)}.b-none{background:var(--none-bg);color:var(--none)}.b-zone{background:var(--accent-soft);color:var(--accent)}
.muted{color:var(--muted)}tr.detail td{background:var(--bg);padding:14px 16px}
.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(280px,1fr));gap:14px}
.box h3{font-size:12px;text-transform:uppercase;letter-spacing:.04em;color:var(--muted);margin:0 0 6px}
.box p{margin:0 0 6px;overflow-wrap:anywhere}.box a{color:var(--accent)}
.inst{width:100%;min-width:0;border:1px solid var(--line)}.inst td,.inst th{padding:5px 8px;font-size:12.5px;position:static;cursor:default}
details summary{cursor:pointer;color:var(--accent)}
footer{max-width:1400px;margin:auto;padding:0 16px 30px;color:var(--muted);font-size:12px}
@media (max-width:640px){header{padding-top:14px}.count{margin-left:0;width:100%}}
</style></head><body>
<header><h1>Indonesia beach places for students (17–25)</h1>
<p class="sub">260 cities/regencies with a real sea beach. Click a column to sort, a row to expand. Safety: 60 largest researched; 200 not researched. Swimmable sand beach = sand (not pebble, rock or mud) with swimming evidenced.</p>
<div class="stats" id="stats"></div></header>
<div class="controls">
<input type="search" id="q" placeholder="Search city, province, institution…" aria-label="Search">
<details class="ms" id="prov"></details><details class="ms" id="saf"></details><details class="ms" id="umk"></details><details class="ms" id="swim"></details><details class="ms" id="zone"></details>
<button id="expand">Expand all</button><button id="collapse">Collapse all</button><button id="reset">Reset</button>
<span class="count" id="count"></span></div>
<main><table><thead><tr id="hdr"></tr></thead><tbody id="body"></tbody></table></main>
<footer>Sources: Kemdiktisaintek 2025 (institutions); OpenStreetMap + web verification (beaches); BPS figures via Databoks/GoodStats mirrors (poverty, unemployment, all ages); governor decrees via news (UMK 2026); local news/police reports 15 Apr–15 Oct 2026 (safety); tourism/news pages (swimmable sand beach). "Unverified" = no sourced value found. Average net wage is not published at regency level.</footer>
<script>
const DATA=__DATA__;
const COLS=[["City/Regency","City / regency","s"],["Province","Province","s"],["Institution count","Institutions","n"],["Offering S1/D3/D4","S1/D3/D4","n"],
["Safety (walking around)","Safety","safety"],["UMK 2026 (IDR/month)","UMK 2026 (IDR)","n"],["Unemployment rate % (all ages)","Unemployment %","n"],["Poverty rate %","Poverty %","n"],["Swimmable sand beach","Swim sand beach","swim"],["Priority zone","Priority zone","s"]];
const WRANK={"Yes":0,"No":1,"Unverified":2,"Not researched":3};
const SRANK={"Low":0,"Medium":1,"High":2,"No rating: <2 dated sources found":3,"Not researched":4};
let sortKey="Institution count",sortDir=-1;const open=new Set();
const $=id=>document.getElementById(id);const esc=s=>String(s??"").replace(/[&<>"]/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
const fmt=(v,k)=>typeof v==="number"?(k.startsWith("UMK")?v.toLocaleString("id-ID"):(k.includes("%")?v.toFixed(2):v)):`<span class="muted">${esc(v)}</span>`;
const badge=v=>`<span class="badge ${SRANK[v]<=2?"b-"+v:"b-none"}">${esc(v.startsWith("No rating")?"No rating":v)}</span>`;
const wbadge=v=>`<span class="badge ${v==="Yes"?"b-High":v==="No"?"b-Low":"b-none"}">${esc(v)}</span>`;
const linkify=s=>esc(s).replace(/https?:\/\/[^\s;)]+/g,u=>`<a href="${u}" target="_blank" rel="noopener">${u.length>70?u.slice(0,67)+"…":u}</a>`);
function init(){buildMS("prov","Province",[...new Set(DATA.map(d=>d.Province))].sort());
buildMS("saf","Safety",Object.keys(SRANK).filter(v=>DATA.some(d=>d["Safety (walking around)"]===v)));
buildMS("umk","Wage basis",[...new Set(DATA.map(d=>d["UMK basis"]))].sort());
buildMS("zone","Priority zone",["Java","Bali","Lampung","Outside"].filter(v=>DATA.some(d=>d["Priority zone"]===v)));
buildMS("swim","Swim beach",Object.keys(WRANK).filter(v=>DATA.some(d=>d["Swimmable sand beach"]===v)));
document.addEventListener("click",e=>document.querySelectorAll("details.ms[open]").forEach(m=>{if(!m.contains(e.target))m.open=false}));
$("hdr").innerHTML=COLS.map(([k,l,t])=>`<th data-k="${esc(k)}" class="${t==="n"?"num":""}">${esc(l)}<span class="arrow"></span></th>`).join("");
$("hdr").onclick=e=>{const th=e.target.closest("th");if(!th)return;const k=th.dataset.k;sortDir=k===sortKey?-sortDir:(["s","safety","swim"].includes(COLS.find(c=>c[0]===k)[2])?1:-1);sortKey=k;render()};
$("q").addEventListener("input",render);
$("expand").onclick=()=>{rows().forEach(d=>open.add(d["City/Regency"]+"|"+d.Province));render()};
$("collapse").onclick=()=>{open.clear();render()};
$("reset").onclick=()=>{$("q").value="";Object.values(OFF).forEach(o=>o.clear());syncMS();open.clear();sortKey="Institution count";sortDir=-1;render()};
$("body").onclick=e=>{const cb=e.target.closest("button.copy");if(cb){copyText(cb.dataset.copy,cb);return}if(e.target.closest("a,details"))return;const tr=e.target.closest("tr.row");if(!tr)return;const k=tr.dataset.k;open.has(k)?open.delete(k):open.add(k);render()};
render()}
function copyText(t,btn){const done=()=>{btn.classList.add("ok");btn.title="Copied";setTimeout(()=>{btn.classList.remove("ok");btn.title="Copy name"},1200)};
const fallback=()=>{const a=document.createElement("textarea");a.value=t;a.style.position="fixed";a.style.opacity="0";document.body.appendChild(a);a.select();
try{document.execCommand("copy");done()}catch(e){}a.remove()};
if(navigator.clipboard&&window.isSecureContext)navigator.clipboard.writeText(t).then(done,fallback);else fallback()}
const OFF={prov:new Set(),saf:new Set(),umk:new Set(),swim:new Set(),zone:new Set()},MSVALS={},MSLABEL={};
function buildMS(id,label,vals){MSVALS[id]=vals;MSLABEL[id]=label;const el=$(id);
el.innerHTML=`<summary></summary><div class="menu"><div class="mbtn"><button type="button" data-all>All</button><button type="button" data-none>None</button></div>`+
vals.map(v=>`<label><input type="checkbox" value="${esc(v)}" checked> ${esc(v)}</label>`).join("")+`</div>`;
el.addEventListener("change",e=>{const v=e.target.value;e.target.checked?OFF[id].delete(v):OFF[id].add(v);syncMS();render()});
el.querySelector("[data-all]").onclick=()=>{OFF[id].clear();syncMS();render()};
el.querySelector("[data-none]").onclick=()=>{vals.forEach(v=>OFF[id].add(v));syncMS();render()};syncMS()}
function syncMS(){for(const id in MSVALS){const el=$(id),n=MSVALS[id].length-OFF[id].size;
el.querySelector("summary").textContent=`${MSLABEL[id]} ${n===MSVALS[id].length?"(all)":n+"/"+MSVALS[id].length}`;el.classList.toggle("active",OFF[id].size>0);
el.querySelectorAll("input").forEach(i=>i.checked=!OFF[id].has(i.value))}}
function rows(){const q=$("q").value.trim().toLowerCase();
return DATA.filter(d=>!OFF.prov.has(d.Province)&&!OFF.saf.has(d["Safety (walking around)"])&&!OFF.umk.has(d["UMK basis"])&&!OFF.swim.has(d["Swimmable sand beach"])&&!OFF.zone.has(d["Priority zone"])&&
(!q||(d["City/Regency"]+" "+d.Province+" "+d.institutions.map(i=>i[0]).join(" ")).toLowerCase().includes(q)))
.sort((a,b)=>{const t=COLS.find(c=>c[0]===sortKey)[2];let x=a[sortKey],y=b[sortKey];
if(t==="safety"){x=SRANK[x];y=SRANK[y]}else if(t==="swim"){x=WRANK[x];y=WRANK[y]}else if(t==="n"){const xn=typeof x==="number",yn=typeof y==="number";if(!xn||!yn)return xn===yn?0:(xn?-1:1)}
else{x=String(x);y=String(y);return x.localeCompare(y)*sortDir}return (x-y)*sortDir})}
function detail(d){const inst=d.institutions.map(i=>`<tr><td>${esc(i[0])}</td><td>${esc(i[1])}</td><td>${esc(i[2])}</td><td>${esc(i[3])}</td><td>${esc(i[4])}</td></tr>`).join("");
return `<div class="grid"><div class="box"><h3>Safety evidence</h3><p>${linkify(d["Safety evidence"])}</p></div>
<div class="box"><h3>Swimmable sand beach</h3><p>${linkify(d["Swim beach evidence"])}</p></div>
<div class="box"><h3>Economy</h3><p>UMK basis: ${esc(d["UMK basis"])}</p><p>Average net wage: ${esc(d["Average net wage (IDR/month, all ages)"])}</p>
<p>Unemployment period: ${esc(d["Unemployment period"])}</p><p>Poverty period: ${esc(d["Poverty period"])}</p></div>
<div class="box"><h3>Sources & notes</h3><p>${linkify(d.Sources)}</p><p class="muted">${esc(d.Notes)}</p></div></div>
<details ${d.institutions.length<=12?"open":""}><summary>${d.institutions.length} institutions</summary>
<table class="inst"><thead><tr><th>Name</th><th>Type</th><th>Ownership</th><th>Levels</th><th>S1/D3/D4</th></tr></thead><tbody>${inst}</tbody></table></details>`}
function render(){const r=rows();
document.querySelectorAll("#hdr th").forEach(th=>{const on=th.dataset.k===sortKey;th.classList.toggle("sorted",on);th.querySelector(".arrow").textContent=on?(sortDir>0?"▲":"▼"):""});
$("body").innerHTML=r.map(d=>{const k=d["City/Regency"]+"|"+d.Province,o=open.has(k);
return `<tr class="row${o?" open":""}" data-k="${esc(k)}"><td><span class="caret">▶</span>${esc(d["City/Regency"])}<button class="copy" data-copy="${esc(d["City/Regency"])}" title="Copy name" aria-label="Copy ${esc(d["City/Regency"])}"><svg viewBox="0 0 24 24" width="14" height="14" aria-hidden="true"><rect x="9" y="9" width="11" height="11" rx="2" fill="none" stroke="currentColor" stroke-width="2"/><path d="M5 15V5a2 2 0 0 1 2-2h8" fill="none" stroke="currentColor" stroke-width="2"/></svg></button></td><td>${esc(d.Province)}</td>
<td class="num">${d["Institution count"]}</td><td class="num">${d["Offering S1/D3/D4"]}</td><td>${badge(d["Safety (walking around)"])}</td>
<td class="num">${fmt(d["UMK 2026 (IDR/month)"],"UMK")}</td><td class="num">${fmt(d["Unemployment rate % (all ages)"],"%")}</td><td class="num">${Number.isInteger(d["Poverty rate %"])&&/decimals truncated/.test(d.Notes)?d["Poverty rate %"]+' <span class="muted" title="Source lists only whole numbers; true value is up to 1 point higher">(whole no.)</span>':fmt(d["Poverty rate %"],"%")}</td><td>${wbadge(d["Swimmable sand beach"])}</td><td>${d["Priority zone"]==="Outside"?'<span class="muted">Outside</span>':`<span class="badge b-zone">${esc(d["Priority zone"])}</span>`}</td></tr>`+
(o?`<tr class="detail"><td colspan="10">${detail(d)}</td></tr>`:"")}).join("");
$("count").textContent=`${r.length} of ${DATA.length} places · ${r.reduce((a,d)=>a+d["Institution count"],0).toLocaleString()} institutions`;
const sc={};DATA.forEach(d=>sc[d["Safety (walking around)"]]=(sc[d["Safety (walking around)"]]||0)+1);
$("stats").innerHTML=[["260","beach places"],[DATA.reduce((a,d)=>a+d["Institution count"],0).toLocaleString(),"institutions"],[sc.Medium||0,"Medium safety"],[sc.Low||0,"Low safety"],[DATA.filter(d=>d["Swimmable sand beach"]==="Yes").length,"swimmable sand beach"],[DATA.filter(d=>d["Priority zone"]!=="Outside").length,"in priority zone"],[DATA.filter(d=>d["UMK basis"]==="UMK").length,"own UMK"]]
.map(([b,s])=>`<div class="stat"><b>${b}</b><span>${s}</span></div>`).join("")}
init();
</script></body></html>"""

if __name__ == "__main__":
    main()
