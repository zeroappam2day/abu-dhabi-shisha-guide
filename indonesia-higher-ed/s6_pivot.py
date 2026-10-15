"""Stage 6: add a real (refreshable) Excel pivot table to the delivered workbook via headless LibreOffice."""
import shutil
import subprocess
import time
import zipfile

from openpyxl import load_workbook

import config as C
from common import gate, read_rows

PORT = 2002


def build_pivot(src, dst):
    import uno
    from com.sun.star.beans import PropertyValue
    from com.sun.star.table import CellAddress

    def prop(name, value):
        p = PropertyValue()
        p.Name, p.Value = name, value
        return p

    office = subprocess.Popen(["soffice", "--headless", "--invisible", "--norestore",
                               f"--accept=socket,host=localhost,port={PORT};urp;"],
                              stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        resolver = uno.getComponentContext().ServiceManager.createInstanceWithContext(
            "com.sun.star.bridge.UnoUrlResolver", uno.getComponentContext())
        for _ in range(60):
            try:
                ctx = resolver.resolve(f"uno:socket,host=localhost,port={PORT};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(1)
        desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(str(src)), "_blank", 0, (prop("Hidden", True),))
        sheets = doc.Sheets
        data = sheets.getByName("Institutions")
        cursor = data.createCursor()
        cursor.gotoEndOfUsedArea(False)
        last_col, last_row = cursor.RangeAddress.EndColumn, cursor.RangeAddress.EndRow
        header = [data.getCellByPosition(c, 0).String for c in range(last_col + 1)]

        sheets.insertNewByName(C.PIVOT_SHEET, sheets.Count)
        target = sheets.getByName(C.PIVOT_SHEET)
        tables = target.getDataPilotTables()
        desc = tables.createDataPilotDescriptor()
        desc.setSourceRange(data.getCellRangeByPosition(0, 0, last_col, last_row).RangeAddress)
        fields = desc.getDataPilotFields()
        orient = "com.sun.star.sheet.DataPilotFieldOrientation"
        for name in C.PIVOT_ROWS:
            fields.getByIndex(header.index(name)).Orientation = uno.Enum(orient, "ROW")
        for name in C.PIVOT_COLUMNS:
            fields.getByIndex(header.index(name)).Orientation = uno.Enum(orient, "COLUMN")
        count = fields.getByIndex(header.index(C.PIVOT_COUNT))
        count.Orientation = uno.Enum(orient, "DATA")
        count.Function = uno.Enum("com.sun.star.sheet.GeneralFunction", "COUNT")
        anchor = CellAddress()
        anchor.Sheet, anchor.Column, anchor.Row = target.RangeAddress.Sheet, 0, 0
        tables.insertNewByName("InstitutionsByCity", anchor, desc)
        doc.storeToURL(uno.systemPathToFileUrl(str(dst)), (prop("FilterName", "Calc MS Excel 2007 XML"),))
        doc.close(True)
    finally:
        office.terminate()
        office.wait(timeout=30)


def main():
    rows = read_rows("s3_institutions.csv")
    draft = C.WORKBOOK.with_suffix(".pivot.xlsx")
    build_pivot(C.WORKBOOK, draft)

    with zipfile.ZipFile(draft) as z:
        has_pivot = any(n.startswith("xl/pivotTables/") for n in z.namelist())
    wb = load_workbook(draft, read_only=True)
    inst = list(wb["Institutions"].iter_rows(min_row=2, values_only=True))
    pivot_cells = [c for row in wb[C.PIVOT_SHEET].iter_rows(values_only=True) for c in row if c is not None]
    grand = max((c for c in pivot_cells if isinstance(c, (int, float))), default=0)
    gate(6, [
        ("pivot table object present (refreshable in Excel)", has_pivot, ""),
        ("Institutions sheet unchanged by the pivot step", len(inst) == len(rows), f"{len(inst)} rows"),
        ("pivot grand total equals institution count", grand == len(rows), f"grand total {grand}"),
    ])
    wb.close()
    shutil.move(draft, C.WORKBOOK)
    print(f"Saved {C.WORKBOOK}")


if __name__ == "__main__":
    main()
